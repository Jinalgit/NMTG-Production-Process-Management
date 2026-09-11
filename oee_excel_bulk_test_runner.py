from __future__ import annotations

import csv
import logging
import math
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl


# ============================================================
# CONFIG
# ============================================================

PROJECT_ROOT = Path(r"D:\Het\demo2")
OEE_ROOT = Path(r"C:\Users\IT system\Desktop\OEE")

FOLDERS = ("A", "B", "C", "D", "E")

ADMIN_ID = 1
ADMIN_USERNAME = "admin"

REPORT_DIR = PROJECT_ROOT / "oee_test_reports"

REJECTION_REASON = "OEE Excel quantity mismatch test"

# User requested automatic execution.
APPLY_CHANGES = True

# None = process every safe match.
MAX_CASES = None

logging.getLogger("mysql.connector").setLevel(logging.WARNING)


# ============================================================
# NORMALIZATION
# ============================================================

def text(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    ).strip()


def key(value):
    return text(value).casefold()


def header(value):
    return text(value).upper()


def jc_key(value):
    value = text(value)

    if not value:
        return ""

    compact = value.replace(" ", "")

    if compact.isdigit():
        return compact.lstrip("0") or "0"

    return re.sub(
        r"[^a-z0-9]+",
        "",
        value.casefold()
    )


def person_key(value):
    return re.sub(
        r"[^a-z0-9]+",
        "",
        text(value).casefold()
    )


def machine_key(value):
    return re.sub(
        r"[^a-z0-9]+",
        "",
        text(value).casefold()
    )


def is_oee_process(value):
    value = key(value)

    return (
        value.startswith("cnc machining")
        or value.startswith("vmc machining")
    )


def category(value):
    value = key(value)

    if value.startswith("cnc machining"):
        return "CNC"

    if value.startswith("vmc machining"):
        return "VMC"

    return ""


def process_match(excel_process, jms_process):
    excel_process = key(excel_process)
    jms_process = key(jms_process)

    if excel_process == jms_process:
        return True

    # Excel can contain generic CNC/VMC while JMS
    # contains 1st Side / 2nd Side / etc.
    if (
        excel_process in (
            "cnc machining",
            "vmc machining",
        )
        and jms_process.startswith(excel_process)
    ):
        return True

    return False


# ============================================================
# VALUE CONVERSION
# ============================================================

def qty(value):
    if value in (None, ""):
        return None

    try:
        value = float(value)
    except Exception:
        return None

    if not math.isfinite(value) or value < 0:
        return None

    rounded = round(value)

    if abs(value - rounded) > 1e-9:
        return None

    return int(rounded)


def nonnegative_float(value):
    if value in (None, ""):
        return 0.0

    try:
        value = float(value)
    except Exception:
        return None

    if not math.isfinite(value) or value < 0:
        return None

    return value


def iso_date(value):
    if value in (None, ""):
        return ""

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    value = text(value)

    for fmt in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%d.%m.%Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(
                value,
                fmt
            ).date().isoformat()
        except ValueError:
            pass

    return ""


def shift_text(value):
    if value is None:
        return ""

    if isinstance(value, (int, float)):
        if float(value).is_integer():
            return str(int(value))

    return text(value)


def time_text(value):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        value = value.time()

    if isinstance(value, time):
        return (
            f"{value.hour:02d}:"
            f"{value.minute:02d}:"
            f"{value.second:02d}"
        )

    if isinstance(value, timedelta):
        total = int(
            round(value.total_seconds())
        ) % 86400

        h, rem = divmod(total, 3600)
        m, s = divmod(rem, 60)

        return f"{h:02d}:{m:02d}:{s:02d}"

    if isinstance(value, (int, float)):
        value = float(value)

        if 0 <= value < 1:
            total = int(
                round(value * 86400)
            ) % 86400

            h, rem = divmod(total, 3600)
            m, s = divmod(rem, 60)

            return f"{h:02d}:{m:02d}:{s:02d}"

    value = text(value)

    for fmt in (
        "%H:%M:%S",
        "%H:%M",
    ):
        try:
            parsed = datetime.strptime(
                value,
                fmt
            ).time()

            return (
                f"{parsed.hour:02d}:"
                f"{parsed.minute:02d}:"
                f"{parsed.second:02d}"
            )
        except ValueError:
            pass

    return None


def duration(value):
    """
    Exact Excel-style input interpretation:

    < 1  -> Excel fraction of day
    >= 1 -> minutes
    """

    if value in (None, ""):
        return 0, 0

    if isinstance(value, datetime):
        value = value.time()

    if isinstance(value, time):
        seconds = (
            value.hour * 3600
            + value.minute * 60
            + value.second
        )

    elif isinstance(value, timedelta):
        seconds = int(
            round(value.total_seconds())
        )

    elif isinstance(value, (int, float)):
        number = float(value)

        if not math.isfinite(number):
            return None

        if number < 0:
            return None

        minutes = (
            number * 1440
            if number < 1
            else number
        )

        seconds = int(
            round(minutes * 60)
        )

    else:
        raw = text(value)

        try:
            parts = raw.split(":")

            if len(parts) == 3:
                h, m, s = map(
                    lambda x: int(float(x)),
                    parts
                )

                seconds = (
                    h * 3600
                    + m * 60
                    + s
                )

            elif len(parts) == 2:
                m, s = map(
                    lambda x: int(float(x)),
                    parts
                )

                seconds = m * 60 + s

            else:
                number = float(raw)

                minutes = (
                    number * 1440
                    if number < 1
                    else number
                )

                seconds = int(
                    round(minutes * 60)
                )

        except Exception:
            return None

    minutes, seconds = divmod(
        seconds,
        60
    )

    return int(minutes), int(seconds)


# ============================================================
# EXCEL HEADER MAP
# ============================================================

ALIASES = {
    "jc": {
        "JC",
        "JC NO",
        "JC NO.",
        "JOB CARD",
        "JOB CARD NO",
        "JOB CARD NO.",
    },

    "production": {
        "PRODUCTION",
        "PRODUCTION QTY",
        "PRODUCTION QUANTITY",
        "OK QTY",
    },

    "rejection": {
        "REJECTION",
        "REJECTION QTY",
        "REJECTED QTY",
        "REJ QTY",
    },

    "rework": {
        "REWORK",
        "REWORK QTY",
        "R/W QTY",
    },

    "process": {
        "OPERATION",
        "PROCESS",
        "PROCESS NAME",
    },

    "machine_no": {
        "M/C NO.",
        "M/C NO",
        "MACHINE NO",
        "MACHINE NO.",
    },

    "machine_name": {
        "MACHINE",
        "MACHINE NAME",
    },

    "operator": {
        "OPERATOR",
        "OPERATOR NAME",
    },

    "date": {
        "DATE",
    },

    "shift": {
        "SHIFT",
    },

    "cycle": {
        "CYCLE TIME",
    },

    "load": {
        "LOAD/UNLOAD",
        "LOAD / UNLOAD",
        "LOAD/UNLOAD TIME",
    },

    "start": {
        "START",
        "START TIME",
    },

    "end": {
        "END",
        "END TIME",
    },
}


def find_header_row(ws):
    for row_no in range(
        1,
        min(ws.max_row, 30) + 1
    ):
        columns = {}

        for col_no in range(
            1,
            ws.max_column + 1
        ):
            value = header(
                ws.cell(
                    row_no,
                    col_no
                ).value
            )

            if value:
                columns[value] = col_no

        has_jc = any(
            x in columns
            for x in ALIASES["jc"]
        )

        has_qty = any(
            x in columns
            for x in ALIASES["production"]
        )

        if has_jc and has_qty:
            return row_no, columns

    return None, {}


def get_col(columns, name):
    for alias in ALIASES[name]:
        if alias in columns:
            return columns[alias]

    return None


def loss_columns(ws, header_row):
    found = {}

    for row_no in range(
        max(1, header_row - 3),
        min(ws.max_row, header_row + 3) + 1
    ):
        for col_no in range(
            1,
            ws.max_column + 1
        ):
            value = header(
                ws.cell(
                    row_no,
                    col_no
                ).value
            )

            match = re.match(
                r"^A(1[0-9]|2[0-7]|[1-9])"
                r"(?:\b|[^0-9])",
                value
            )

            if match:
                code = "A" + str(
                    int(match.group(1))
                )

                found.setdefault(
                    code,
                    col_no
                )

    # Existing OEE workbooks:
    # AH:BH = A1:A27
    if len(found) != 27:
        if ws.max_column >= 60:
            found = {
                f"A{i}": 33 + i
                for i in range(1, 28)
            }

    return found


# ============================================================
# READ ALL A-E EXCEL FILES
# ============================================================

def scan_excel():
    rows = []

    scanned_files = 0

    for folder_name in FOLDERS:
        folder = OEE_ROOT / folder_name

        if not folder.exists():
            print(
                "WARNING - folder missing:",
                folder
            )
            continue

        files = sorted([
            p
            for p in folder.rglob("*")
            if (
                p.is_file()
                and p.suffix.lower()
                in (".xlsx", ".xlsm")
                and not p.name.startswith("~$")
            )
        ])

        for path in files:
            try:
                wb = openpyxl.load_workbook(
                    path,
                    read_only=True,
                    data_only=True
                )
            except Exception as exc:
                print(
                    "WARNING - cannot open:",
                    path.name,
                    exc
                )
                continue

            scanned_files += 1

            for ws in wb.worksheets:
                header_row, headers = (
                    find_header_row(ws)
                )

                if not header_row:
                    continue

                cols = {
                    name: get_col(
                        headers,
                        name
                    )
                    for name in ALIASES
                }

                if (
                    not cols["jc"]
                    or not cols["production"]
                ):
                    continue

                loss_cols = loss_columns(
                    ws,
                    header_row
                )

                for row_no in range(
                    header_row + 1,
                    ws.max_row + 1
                ):
                    raw_jc = ws.cell(
                        row_no,
                        cols["jc"]
                    ).value

                    if raw_jc in (None, ""):
                        continue

                    production = qty(
                        ws.cell(
                            row_no,
                            cols["production"]
                        ).value
                    )

                    if production is None:
                        continue

                    process = (
                        ws.cell(
                            row_no,
                            cols["process"]
                        ).value
                        if cols["process"]
                        else ""
                    )

                    if not is_oee_process(process):
                        continue

                    losses = []
                    loss_error = ""

                    if len(loss_cols) != 27:
                        loss_error = (
                            "LOSS_COLUMNS_NOT_FOUND"
                        )

                    else:
                        for i in range(
                            1,
                            28
                        ):
                            code = f"A{i}"

                            value = (
                                nonnegative_float(
                                    ws.cell(
                                        row_no,
                                        loss_cols[code]
                                    ).value
                                )
                            )

                            if value is None:
                                loss_error = (
                                    f"INVALID_{code}"
                                )
                                break

                            losses.append({
                                "loss_code": code,
                                "loss_minutes": value
                            })

                    def cell(name):
                        col = cols.get(name)

                        if not col:
                            return None

                        return ws.cell(
                            row_no,
                            col
                        ).value

                    rejection = (
                        qty(cell("rejection"))
                        if cols["rejection"]
                        else 0
                    )

                    rework = (
                        qty(cell("rework"))
                        if cols["rework"]
                        else 0
                    )

                    rejection = (
                        rejection
                        if rejection is not None
                        else 0
                    )

                    rework = (
                        rework
                        if rework is not None
                        else 0
                    )

                    rows.append({
                        "folder": folder_name,
                        "file": path.name,
                        "sheet": ws.title,
                        "row": row_no,

                        "jc": text(raw_jc),
                        "jc_key": jc_key(raw_jc),

                        "production": production,

                        "excel_rejection": rejection,
                        "excel_rework": rework,

                        "process": text(process),

                        "operator": text(
                            cell("operator")
                        ),

                        "machine_no": text(
                            cell("machine_no")
                        ),

                        "machine_name": text(
                            cell("machine_name")
                        ),

                        "date": iso_date(
                            cell("date")
                        ),

                        "shift": shift_text(
                            cell("shift")
                        ),

                        "start": time_text(
                            cell("start")
                        ),

                        "end": time_text(
                            cell("end")
                        ),

                        "cycle": duration(
                            cell("cycle")
                        ),

                        "load": duration(
                            cell("load")
                        ),

                        "losses": losses,
                        "loss_error": loss_error,
                    })

            wb.close()

    print(
        "Excel files scanned:",
        scanned_files
    )

    print(
        "CNC/VMC Excel rows:",
        len(rows)
    )

    return rows


# ============================================================
# FLASK SESSION HELPERS
# ============================================================

def set_session(client, user):
    with client.session_transaction() as sess:
        sess.clear()

        sess["role"] = user["role"]
        sess["user_id"] = int(
            user["id"]
        )

        sess["username"] = (
            user.get("username")
            or ""
        )

        sess["full_name"] = (
            user.get("full_name")
            or ""
        )


def admin():
    return {
        "id": ADMIN_ID,
        "username": ADMIN_USERNAME,
        "full_name": "Admin",
        "role": "admin",
    }


def response_json(response):
    data = response.get_json(
        silent=True
    )

    return (
        data
        if isinstance(data, dict)
        else {}
    )


def list_rows(data):
    for name in (
        "job_cards",
        "cards",
        "data",
        "items",
    ):
        value = data.get(name)

        if isinstance(value, list):
            return [
                row
                for row in value
                if isinstance(row, dict)
            ]

    return []


# ============================================================
# GET ALL ACTIVE OPERATORS
# ============================================================

def get_operators(app):
    client = app.test_client()

    set_session(
        client,
        admin()
    )

    response = client.get(
        "/api/users"
    )

    data = response_json(
        response
    )

    if (
        response.status_code != 200
        or not data.get("success")
    ):
        raise RuntimeError(
            "Unable to read /api/users: "
            + str(data)
        )

    operators = []

    for user in (
        data.get("users")
        or data.get("data")
        or []
    ):
        if (
            key(user.get("role"))
            != "operator"
        ):
            continue

        if (
            int(
                user.get("is_active")
                or 0
            )
            != 1
        ):
            continue

        operators.append({
            "id": int(user["id"]),
            "username": text(
                user.get("username")
            ),
            "full_name": text(
                user.get("full_name")
            ),
            "role": "operator",
        })

    return operators


# ============================================================
# GET ALL LIVE CNC/VMC OPERATOR CARDS
# ============================================================

def get_live_cards(app, operators):
    cards = []

    for operator in operators:
        client = app.test_client()

        set_session(
            client,
            operator
        )

        response = client.get(
            "/api/operator/job_cards"
        )

        data = response_json(
            response
        )

        if (
            response.status_code != 200
            or not data.get("success")
        ):
            continue

        for card in list_rows(data):
            process = (
                card.get("wip_status")
                or card.get(
                    "current_process"
                )
                or card.get(
                    "process_name"
                )
                or ""
            )

            if not is_oee_process(
                process
            ):
                continue

            row = dict(card)

            row["_operator"] = (
                operator
            )

            row["_process"] = text(
                process
            )

            row["_next"] = text(
                card.get(
                    "next_process"
                )
                or "Store"
            )

            row["_jc_key"] = jc_key(
                card.get(
                    "job_card_no"
                )
            )

            row["_item_key"] = key(
                card.get(
                    "item_name"
                )
            )

            cards.append(row)

    return cards


# ============================================================
# MACHINE MASTER
# ============================================================

def get_machines(app):
    client = app.test_client()

    set_session(
        client,
        admin()
    )

    response = client.get(
        "/api/oee/master-data"
    )

    data = response_json(
        response
    )

    if (
        response.status_code != 200
        or not data.get("success")
    ):
        raise RuntimeError(
            "Unable to read OEE machines: "
            + str(data)
        )

    return data.get(
        "machines"
    ) or []


def match_machine(
    excel_row,
    current_process,
    machines
):
    cat = category(
        current_process
    )

    eligible = [
        machine
        for machine in machines
        if text(
            machine.get(
                "machine_category"
            )
        ).upper() == cat
    ]

    no_key = machine_key(
        excel_row["machine_no"]
    )

    name_key = machine_key(
        excel_row["machine_name"]
    )

    matches = []

    if no_key:
        matches = [
            machine
            for machine in eligible
            if machine_key(
                machine.get(
                    "machine_no"
                )
            ) == no_key
        ]

    if not matches and name_key:
        matches = [
            machine
            for machine in eligible
            if machine_key(
                machine.get(
                    "machine_name"
                )
            ) == name_key
        ]

    if len(matches) == 1:
        return matches[0]

    return None


# ============================================================
# CURRENT STAGE AVAILABLE QTY
# ============================================================

def available_qty(card):
    job_qty = (
        qty(
            card.get(
                "job_card_qty"
            )
        )
        or 0
    )

    so_qty = (
        qty(
            card.get(
                "so_qty"
            )
        )
        or 0
    )

    actual = (
        qty(
            card.get(
                "actual_qty"
            )
        )
        or 0
    )

    rejected = (
        qty(
            card.get(
                "rejected_qty"
            )
        )
        or 0
    )

    hold = (
        qty(
            card.get(
                "hold_qty"
            )
        )
        or 0
    )

    pending = (
        qty(
            card.get(
                "pending_qty"
            )
        )
        or 0
    )

    partial = (
        pending > 0
        or hold > 0
    )

    if partial:
        return (
            actual
            + rejected
            + hold
            + pending,
            True
        )

    prior_outcome = (
        actual > 0
        or rejected > 0
    )

    base = (
        job_qty
        if job_qty > 0
        else so_qty
    )

    return (
        actual
        if prior_outcome
        else base,
        False
    )


# ============================================================
# CHOOSE EXCEL/JMS TEST CASES
# ============================================================

def choose_cases(
    excel_rows,
    live_cards,
    machines
):
    excel_by_jc = defaultdict(
        list
    )

    live_by_jc = defaultdict(
        list
    )

    for row in excel_rows:
        excel_by_jc[
            row["jc_key"]
        ].append(row)

    for card in live_cards:
        live_by_jc[
            card["_jc_key"]
        ].append(card)

    selected = []
    skipped = []

    for jckey, excel_group in (
        excel_by_jc.items()
    ):
        live_group = (
            live_by_jc.get(
                jckey,
                []
            )
        )

        if not live_group:
            for row in excel_group:
                skipped.append({
                    **row,
                    "status": "SKIP",
                    "reason":
                        "JC_NOT_IN_LIVE_CNC_VMC_QUEUE"
                })

            continue

        item_keys = {
            card["_item_key"]
            for card in live_group
        }

        if len(item_keys) != 1:
            for row in excel_group:
                skipped.append({
                    **row,
                    "status": "SKIP",
                    "reason":
                        "MULTIPLE_LIVE_ITEMS_FOR_JC"
                })

            continue

        current_process = (
            live_group[0][
                "_process"
            ]
        )

        matches = [
            row
            for row in excel_group
            if process_match(
                row["process"],
                current_process
            )
        ]

        if not matches:
            for row in excel_group:
                skipped.append({
                    **row,
                    "status": "SKIP",
                    "reason":
                        "CURRENT_STAGE_MISMATCH"
                })

            continue

        # Prefer latest date, then latest row.
        matches.sort(
            key=lambda row: (
                row["date"],
                row["file"],
                row["sheet"],
                row["row"],
            ),
            reverse=True
        )

        excel_row = matches[0]

        # Prefer Excel operator if it exists
        # as one of the eligible JMS operators.
        excel_operator = person_key(
            excel_row["operator"]
        )

        selected_card = None

        if excel_operator:
            for card in live_group:
                op = card[
                    "_operator"
                ]

                if excel_operator in {
                    person_key(
                        op.get(
                            "username"
                        )
                    ),
                    person_key(
                        op.get(
                            "full_name"
                        )
                    ),
                }:
                    selected_card = card
                    break

        if selected_card is None:
            selected_card = sorted(
                live_group,
                key=lambda card:
                    card[
                        "_operator"
                    ]["id"]
            )[0]

        machine = match_machine(
            excel_row,
            current_process,
            machines
        )

        if machine is None:
            skipped.append({
                **excel_row,
                "status": "SKIP",
                "reason":
                    "MACHINE_NOT_MATCHED"
            })

            continue

        selected.append({
            **excel_row,

            "_card":
                selected_card,

            "_operator":
                selected_card[
                    "_operator"
                ],

            "_machine":
                machine,
        })

        for duplicate in matches[1:]:
            skipped.append({
                **duplicate,
                "status": "SKIP",
                "reason":
                    "DUPLICATE_EXCEL_ROW"
            })

    if MAX_CASES is not None:
        overflow = selected[
            MAX_CASES:
        ]

        selected = selected[
            :MAX_CASES
        ]

        for row in overflow:
            skipped.append({
                **row,
                "status": "SKIP",
                "reason":
                    "MAX_CASES_LIMIT"
            })

    return selected, skipped


# ============================================================
# REFRESH CARD BEFORE SUBMISSION
# ============================================================

def refresh_card(
    app,
    operator,
    jckey,
    item_key
):
    client = app.test_client()

    set_session(
        client,
        operator
    )

    response = client.get(
        "/api/operator/job_cards"
    )

    data = response_json(
        response
    )

    if (
        response.status_code != 200
        or not data.get("success")
    ):
        return None

    for card in list_rows(data):
        if (
            jc_key(
                card.get(
                    "job_card_no"
                )
            )
            != jckey
        ):
            continue

        if (
            item_key
            and key(
                card.get(
                    "item_name"
                )
            )
            != item_key
        ):
            continue

        return card

    return None


# ============================================================
# RUN ONE TEST CASE
# ============================================================

def run_case(
    app,
    row
):
    operator = row[
        "_operator"
    ]

    old_card = row[
        "_card"
    ]

    item_key = key(
        old_card.get(
            "item_name"
        )
    )

    card = refresh_card(
        app,
        operator,
        row["jc_key"],
        item_key
    )

    if card is None:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "LIVE_CARD_CHANGED"
        }

    current_process = text(
        card.get(
            "wip_status"
        )
        or card.get(
            "current_process"
        )
        or ""
    )

    if not process_match(
        row["process"],
        current_process
    ):
        return {
            **row,
            "status": "SKIP",
            "reason":
                "STAGE_CHANGED_BEFORE_SUBMIT"
        }

    stage_qty, partial = (
        available_qty(card)
    )

    if partial:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "CURRENT_STAGE_IS_PARTIAL"
        }

    if stage_qty <= 0:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "JMS_STAGE_QTY_ZERO"
        }

    excel_ok = row[
        "production"
    ]

    if excel_ok > stage_qty:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "EXCEL_QTY_GREATER_THAN_JMS",
            "jms_qty": stage_qty,
        }

    # THIS IS THE TEST RULE:
    #
    # JMS available = 40
    # Excel production = 36
    #
    # OK       = 36
    # Rejected = 4
    #
    rejected = (
        stage_qty
        - excel_ok
    )

    if not row["date"]:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "EXCEL_DATE_MISSING"
        }

    if not row["shift"]:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "EXCEL_SHIFT_MISSING"
        }

    if row["cycle"] is None:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "INVALID_CYCLE_TIME"
        }

    if row["load"] is None:
        return {
            **row,
            "status": "SKIP",
            "reason":
                "INVALID_LOAD_UNLOAD"
        }

    if row["loss_error"]:
        return {
            **row,
            "status": "SKIP",
            "reason":
                row["loss_error"]
        }

    cycle_min, cycle_sec = (
        row["cycle"]
    )

    load_min, load_sec = (
        row["load"]
    )

    next_process = text(
        card.get(
            "next_process"
        )
        or "Store"
    )

    source = (
        f"{row['folder']}\\"
        f"{row['file']} | "
        f"{row['sheet']} "
        f"row {row['row']}"
    )

    payload = {
        "job_card_no":
            text(
                card.get(
                    "job_card_no"
                )
            ),

        "item_name":
            text(
                card.get(
                    "item_name"
                )
            ),

        "new_stage":
            next_process,

        "changed_by":
            (
                operator.get(
                    "full_name"
                )
                or operator.get(
                    "username"
                )
                or "Operator"
            ),

        "actual_qty":
            excel_ok,

        "ok_qty":
            excel_ok,

        "rejected_qty":
            rejected,

        "hold_qty":
            0,

        "rejection_reason":
            (
                REJECTION_REASON
                if rejected > 0
                else ""
            ),

        "hold_reason":
            "",

        "rework_qty":
            0,

        "rework_remarks":
            "",

        "stage_remark":
            "AUTO OEE TEST | "
            + source,

        "oee": {
            "entry_date":
                row["date"],

            "shift_name":
                row["shift"],

            "machine_id":
                int(
                    row[
                        "_machine"
                    ]["id"]
                ),

            "start_time":
                row["start"],

            "end_time":
                row["end"],

            "cycle_minutes":
                cycle_min,

            "cycle_seconds":
                cycle_sec,

            "load_unload_minutes":
                load_min,

            "load_unload_seconds":
                load_sec,

            "losses":
                row["losses"],
        },
    }

    if not APPLY_CHANGES:
        return {
            **row,
            "status": "DRY_RUN",
            "reason":
                "READY",
            "jms_qty":
                stage_qty,
            "submitted_ok":
                excel_ok,
            "submitted_rejected":
                rejected,
        }

    client = app.test_client()

    set_session(
        client,
        operator
    )

    response = client.post(
        "/api/wip/update",
        json=payload
    )

    data = response_json(
        response
    )

    if (
        response.status_code != 200
        or not data.get("success")
    ):
        return {
            **row,
            "status": "FAIL",
            "reason":
                "WIP_UPDATE_FAILED: "
                + text(
                    data.get(
                        "error"
                    )
                    or response.status_code
                ),
            "jms_qty":
                stage_qty,
            "submitted_ok":
                excel_ok,
            "submitted_rejected":
                rejected,
        }

    completion_id = data.get(
        "operator_completion_id"
    )

    oee_id = data.get(
        "oee_entry_id"
    )

    if not oee_id:
        return {
            **row,
            "status": "FAIL",
            "reason":
                "OEE_ID_MISSING",
            "completion_id":
                completion_id,
        }

    # Verify through the same results API
    # used by the OEE frontend.
    response = client.get(
        "/api/oee/results",
        query_string={
            "search":
                payload[
                    "job_card_no"
                ],

            "per_page":
                200,
        }
    )

    result_data = response_json(
        response
    )

    if (
        response.status_code != 200
        or not result_data.get(
            "success"
        )
    ):
        return {
            **row,
            "status": "FAIL",
            "reason":
                "RESULTS_API_FAILED",
            "oee_entry_id":
                oee_id,
        }

    stored = next(
        (
            item
            for item in (
                result_data.get(
                    "data"
                )
                or []
            )
            if int(
                item.get(
                    "oee_entry_id"
                )
                or 0
            )
            == int(oee_id)
        ),
        None
    )

    if stored is None:
        return {
            **row,
            "status": "FAIL",
            "reason":
                "ACTIVE_OEE_NOT_FOUND",
            "oee_entry_id":
                oee_id,
        }

    return {
        **row,

        "status":
            "PASS",

        "reason":
            "OEE_CREATED_AND_VERIFIED",

        "jms_job_card":
            payload[
                "job_card_no"
            ],

        "item_name":
            payload[
                "item_name"
            ],

        "jms_process":
            current_process,

        "next_process":
            next_process,

        "operator_used":
            (
                operator.get(
                    "full_name"
                )
                or operator.get(
                    "username"
                )
            ),

        "machine_used":
            (
                row[
                    "_machine"
                ].get(
                    "machine_no"
                )
                or row[
                    "_machine"
                ].get(
                    "machine_name"
                )
            ),

        "jms_qty":
            stage_qty,

        "submitted_ok":
            excel_ok,

        "submitted_rejected":
            rejected,

        "completion_id":
            completion_id,

        "oee_entry_id":
            oee_id,

        "AR":
            stored.get(
                "detailed_ar_ratio"
            ),

        "PR":
            stored.get(
                "detailed_pr_ratio"
            ),

        "QR":
            stored.get(
                "detailed_qr_ratio"
            ),

        "OEE":
            stored.get(
                "detailed_oee_ratio"
            ),
    }


# ============================================================
# REPORT
# ============================================================

def write_report(results):
    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    path = REPORT_DIR / (
        "OEE_Bulk_Test_Report_"
        + datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
        + ".csv"
    )

    fields = [
        "status",
        "reason",
        "folder",
        "file",
        "sheet",
        "row",
        "jc",
        "jms_job_card",
        "item_name",
        "process",
        "jms_process",
        "next_process",
        "operator",
        "operator_used",
        "machine_no",
        "machine_used",
        "jms_qty",
        "production",
        "submitted_ok",
        "submitted_rejected",
        "completion_id",
        "oee_entry_id",
        "AR",
        "PR",
        "QR",
        "OEE",
    ]

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=fields
        )

        writer.writeheader()

        for result in results:
            writer.writerow({
                field:
                    result.get(
                        field,
                        ""
                    )
                for field in fields
            })

    return path


# ============================================================
# MAIN
# ============================================================

def main():
    if not PROJECT_ROOT.exists():
        raise SystemExit(
            "FAIL - project folder missing: "
            + str(PROJECT_ROOT)
        )

    if not OEE_ROOT.exists():
        raise SystemExit(
            "FAIL - OEE root missing: "
            + str(OEE_ROOT)
        )

    os.chdir(
        PROJECT_ROOT
    )

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(
            0,
            str(PROJECT_ROOT)
        )

    from app import app

    print()
    print("=" * 72)
    print(
        "NMTG JMS - AUTOMATIC EXCEL -> OEE TEST"
    )
    print("=" * 72)

    print()
    print(
        "[1] Scanning Excel folders A-E..."
    )

    excel_rows = scan_excel()

    print()
    print(
        "[2] Reading active JMS operators..."
    )

    operators = get_operators(
        app
    )

    print(
        "Active operators:",
        len(operators)
    )

    print()
    print(
        "[3] Reading live CNC/VMC queues..."
    )

    live_cards = get_live_cards(
        app,
        operators
    )

    print(
        "Live CNC/VMC cards:",
        len(live_cards)
    )

    print()
    print(
        "[4] Reading OEE machine master..."
    )

    machines = get_machines(
        app
    )

    print(
        "Machines:",
        len(machines)
    )

    print()
    print(
        "[5] Matching Excel rows "
        "against live JMS..."
    )

    selected, skipped = (
        choose_cases(
            excel_rows,
            live_cards,
            machines
        )
    )

    print(
        "Safe test cases:",
        len(selected)
    )

    print()
    print(
        "[6] Running automatic OEE tests..."
    )

    results = list(
        skipped
    )

    for index, row in enumerate(
        selected,
        start=1
    ):
        print(
            f"{index}/{len(selected)} "
            f"| JC {row['jc']} "
            f"| {row['process']} "
            f"| Excel Qty "
            f"{row['production']}"
        )

        try:
            result = run_case(
                app,
                row
            )

        except Exception as exc:
            result = {
                **row,
                "status": "FAIL",
                "reason":
                    "UNHANDLED ERROR: "
                    + str(exc)
            }

        results.append(
            result
        )

        print(
            "   ",
            result["status"],
            "-",
            result["reason"]
        )

        if result["status"] == "PASS":
            print(
                "    OEE ID:",
                result.get(
                    "oee_entry_id"
                ),
                "| OEE:",
                result.get(
                    "OEE"
                )
            )

    report = write_report(
        results
    )

    counts = Counter(
        row["status"]
        for row in results
    )

    reasons = Counter(
        row["reason"].split(":")[0]
        for row in results
        if row["status"] != "PASS"
    )

    print()
    print("=" * 72)
    print(
        "FINAL RESULT"
    )
    print("=" * 72)

    print(
        "Excel CNC/VMC rows :",
        len(excel_rows)
    )

    print(
        "Selected tests     :",
        len(selected)
    )

    print(
        "PASS               :",
        counts.get(
            "PASS",
            0
        )
    )

    print(
        "FAIL               :",
        counts.get(
            "FAIL",
            0
        )
    )

    print(
        "SKIP               :",
        counts.get(
            "SKIP",
            0
        )
    )

    print(
        "Report             :",
        report
    )

    if counts.get("PASS", 0):
        print()
        print(
            "PASS OEE records remain ACTIVE."
        )

        print(
            "They should now be visible "
            "on the standalone OEE page."
        )

    if reasons:
        print()
        print(
            "Main skipped/failed reasons:"
        )

        for reason, count in (
            reasons.most_common(
                10
            )
        ):
            print(
                f"{count:>4} - "
                f"{reason}"
            )

    print()
    print(
        "NO DIRECT MYSQL WRITE "
        "WAS USED."
    )

    print("=" * 72)


if __name__ == "__main__":
    main()