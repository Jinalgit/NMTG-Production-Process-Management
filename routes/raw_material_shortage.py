import re
import time
from io import BytesIO
from datetime import date
from decimal import Decimal, InvalidOperation

from flask import Blueprint, render_template, session, redirect, url_for, send_file

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from db import get_connection
from permission_utils import is_gaurang_special_user


raw_material_shortage_bp = Blueprint(
    "raw_material_shortage",
    __name__,
)


# RAW_MATERIAL_SHORTAGE_BULK_V2


def _clean(value):
    return str(value or "").strip()


def _material_key(value):
    return re.sub(
        r"\s+",
        "",
        _clean(value).lower(),
    )


def _is_forging_ring(material):
    material = _clean(material).lower()

    return (
        "forged" in material
        and "ring" in material
    ) or (
        "forging" in material
        and "ring" in material
    )


def _number_from_cut_size(value):
    match = re.search(
        r"[0-9]+(?:\.[0-9]+)?",
        _clean(value),
    )

    if not match:
        return None

    try:
        return Decimal(match.group(0))
    except InvalidOperation:
        return None


def _decimal_qty(value):
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def _format_shortage_qty(value):
    try:
        value = Decimal(str(value or 0))
    except Exception:
        return str(value or "0")

    if value == value.to_integral():
        return f"{int(value):,}"

    formatted = f"{value:,.2f}"
    return formatted.rstrip("0").rstrip(".")


def _load_raw_material_shortage():
    started = time.perf_counter()

    conn = None
    cursor = None

    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        # ------------------------------------------------------------
        # 1. Fetch current Raw Material-pending JCs ONCE
        # ------------------------------------------------------------

        cursor.execute("""
            SELECT
                ji.job_card_no,
                jc.child_code,
                ji.so_qty AS qty,
                NULLIF(TRIM(ji.material), '') AS jc_material,
                NULLIF(TRIM(ji.cutting_size), '') AS jc_cut_size,

                MAX(jpd.in_time) AS rm_in_time,

                GREATEST(
                    DATEDIFF(
                        CURDATE(),
                        DATE(MAX(jpd.in_time))
                    ),
                    0
                ) AS days_in_stage

            FROM job_card_items ji

            JOIN job_cards jc
                ON BINARY jc.job_card_no
                 = BINARY ji.job_card_no

            JOIN job_card_process_days jpd
                ON BINARY jpd.job_card_no
                 = BINARY ji.job_card_no
               AND LOWER(TRIM(jpd.process_name))
                   = 'raw material'
               AND jpd.in_time IS NOT NULL
               AND jpd.out_time IS NULL
               AND COALESCE(jpd.is_completed, 0) = 0

            WHERE COALESCE(ji.is_deleted, 0) = 0
              AND LOWER(TRIM(ji.wip_status))
                  = 'raw material'

              AND ji.job_card_no <> '0000112671'

            GROUP BY
                ji.id,
                ji.job_card_no,
                jc.child_code,
                ji.so_qty,
                ji.material,
                ji.cutting_size
        """)

        pending_rows = cursor.fetchall() or []

        if not pending_rows:
            return []

        # ------------------------------------------------------------
        # 2. Collect only child codes that need historical BOM fallback
        # ------------------------------------------------------------

        fallback_codes = []

        for row in pending_rows:
            if (
                not _clean(row.get("jc_material"))
                or not _clean(row.get("jc_cut_size"))
            ):
                child_code = _clean(
                    row.get("child_code")
                )

                if child_code:
                    fallback_codes.append(
                        child_code
                    )

        fallback_codes = list(
            dict.fromkeys(fallback_codes)
        )

        bom_by_parent = {}

        # ------------------------------------------------------------
        # 3. Fetch all relevant BOM candidates together
        # ------------------------------------------------------------

        if fallback_codes:
            placeholders = ",".join(
                ["%s"] * len(fallback_codes)
            )

            cursor.execute(
                f"""
                SELECT
                    bl.id,
                    bl.parent_code,
                    bl.child_code,
                    NULLIF(
                        TRIM(bl.cutting_size),
                        ''
                    ) AS cutting_size,
                    NULLIF(
                        TRIM(i.item_description),
                        ''
                    ) AS material

                FROM bom_links bl

                LEFT JOIN items i
                    ON BINARY i.item_code
                     = BINARY bl.child_code

                WHERE bl.make_buy = 'I'
                  AND COALESCE(
                        bl.is_alternate,
                        0
                      ) = 0

                  AND BINARY bl.parent_code IN (
                      {placeholders}
                  )

                ORDER BY
                    bl.parent_code,
                    bl.id
                """,
                tuple(fallback_codes),
            )

            for bom in cursor.fetchall() or []:
                parent_code = _clean(
                    bom.get("parent_code")
                )

                bom_by_parent.setdefault(
                    parent_code,
                    [],
                ).append(bom)

        # ------------------------------------------------------------
        # 4. Resolve effective source + calculate in Python
        # ------------------------------------------------------------

        grouped = {}

        for row in pending_rows:

            child_code = _clean(
                row.get("child_code")
            )

            candidates = bom_by_parent.get(
                child_code,
                [],
            )

            jc_material = _clean(
                row.get("jc_material")
            )

            jc_cut_size = _clean(
                row.get("jc_cut_size")
            )

            # Material priority:
            # JC/PDF first -> old BOM fallback
            effective_material = jc_material

            if not effective_material:
                for candidate in candidates:
                    material = _clean(
                        candidate.get("material")
                    )

                    if material:
                        effective_material = material
                        break

            if not effective_material:
                continue

            # Cut Size priority:
            # JC/PDF first -> matching BOM -> first usable BOM
            effective_cut_size = jc_cut_size

            if (
                not effective_cut_size
                and not _is_forging_ring(
                    effective_material
                )
            ):
                material_match_key = _material_key(
                    effective_material
                )

                for candidate in candidates:
                    candidate_material = _clean(
                        candidate.get("material")
                    )

                    candidate_cut = _clean(
                        candidate.get("cutting_size")
                    )

                    if (
                        candidate_cut
                        and _material_key(
                            candidate_material
                        ) == material_match_key
                    ):
                        effective_cut_size = (
                            candidate_cut
                        )
                        break

                if not effective_cut_size:
                    for candidate in candidates:
                        candidate_cut = _clean(
                            candidate.get(
                                "cutting_size"
                            )
                        )

                        if candidate_cut:
                            effective_cut_size = (
                                candidate_cut
                            )
                            break

            qty = _decimal_qty(
                row.get("qty")
            )

            if _is_forging_ring(
                effective_material
            ):
                required_qty = qty
                uom = "No"

            else:
                cut_size = _number_from_cut_size(
                    effective_cut_size
                )

                if cut_size is None:
                    continue

                required_qty = (
                    cut_size + Decimal("2")
                ) * qty

                uom = "mm"

            key = _material_key(
                effective_material
            )

            days = int(
                row.get("days_in_stage")
                or 0
            )

            if key not in grouped:
                grouped[key] = {
                    "material": effective_material,
                    "uom": uom,
                    "shortage_qty": Decimal("0"),
                    "days_in_stage": days,
                }

            grouped[key]["shortage_qty"] += (
                required_qty
            )

            grouped[key]["days_in_stage"] = max(
                grouped[key]["days_in_stage"],
                days,
            )

        rows = list(grouped.values())

        rows.sort(
            key=lambda row: (
                -row["days_in_stage"],
                row["material"].lower(),
            )
        )

        for row in rows:
            row["shortage_display"] = (
                _format_shortage_qty(
                    row["shortage_qty"]
                )
            )

        elapsed = time.perf_counter() - started

        print(
            "[Raw Material Shortage] "
            f"Pending={len(pending_rows)} "
            f"BOMCodes={len(fallback_codes)} "
            f"Materials={len(rows)} "
            f"Time={elapsed:.3f}s"
        )

        return rows

    finally:
        if cursor:
            cursor.close()

        if conn:
            conn.close()


# RAW_MATERIAL_SHORTAGE_EXCEL_V1
@raw_material_shortage_bp.route(
    "/raw-material-shortage/export-excel"
)
def export_excel():
    role = _clean(
        session.get("role")
    ).lower()

    is_gaurang_special = (
        is_gaurang_special_user()
    )

    if (
        role != "admin"
        and not is_gaurang_special
    ):
        return redirect(
            url_for("pages.index")
        )

    shortage_rows = (
        _load_raw_material_shortage()
    )

    shortage_date = date.today().strftime(
        "%d.%m.%y"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Material Shortage"

    headers = [
        "Raw Material-Pending",
        "UOM",
        f"Shortage {shortage_date}",
        "Days in Stage",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in shortage_rows:
        shortage_qty = row.get(
            "shortage_qty"
        ) or 0

        try:
            shortage_qty = int(
                round(float(shortage_qty))
            )
        except Exception:
            shortage_qty = 0

        ws.append([
            row.get("material") or "",
            row.get("uom") or "",
            shortage_qty,
            int(
                row.get("days_in_stage")
                or 0
            ),
        ])

    for row in ws.iter_rows(
        min_row=2,
        max_col=4,
    ):
        row[1].alignment = Alignment(
            horizontal="center"
        )

        row[2].alignment = Alignment(
            horizontal="right"
        )

        row[2].number_format = (
            '0'
        )

        row[3].alignment = Alignment(
            horizontal="center"
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 18

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = (
        "Raw_Material_Shortage_"
        + date.today().strftime("%Y-%m-%d")
        + ".xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )


@raw_material_shortage_bp.route(
    "/raw-material-shortage"
)
def index():
    role = _clean(
        session.get("role")
    ).lower()

    is_gaurang_special = (
        is_gaurang_special_user()
    )

    if (
        role != "admin"
        and not is_gaurang_special
    ):
        return redirect(
            url_for("pages.index")
        )

    shortage_rows = (
        _load_raw_material_shortage()
    )

    return render_template(
        "raw_material_shortage.html",
        active_page="raw_material_shortage",
        is_gaurang_special=is_gaurang_special,
        shortage_rows=shortage_rows,
        shortage_date=date.today().strftime(
            "%d.%m.%y"
        ),
    )
